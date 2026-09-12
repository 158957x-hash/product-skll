#!/usr/bin/env python3
import argparse,json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
def set_font(run,size=10.5,bold=False):
    run.font.name="微软雅黑"; run._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑"); run.font.size=Pt(size); run.bold=bold
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("spec"); ap.add_argument("output"); a=ap.parse_args(); spec=json.loads(Path(a.spec).read_text(encoding="utf-8-sig")); doc=Document()
    for name in ["Normal","Title","Heading 1","Heading 2","Heading 3"]:
        st=doc.styles[name]; st.font.name="微软雅黑"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑")
    p=doc.add_paragraph();p.alignment=1;r=p.add_run(spec.get("title","系统操作手册"));set_font(r,22,True);doc.add_page_break()
    for sec in spec.get("sections",[]):
        doc.add_heading(sec.get("heading","未命名章节"),level=int(sec.get("level",1)))
        for text in sec.get("paragraphs",[]): doc.add_paragraph(text)
        for step in sec.get("steps",[]):
            value=step if isinstance(step,str) else step.get("text",""); doc.add_paragraph(value,style="List Number")
            if isinstance(step,dict) and step.get("image") and Path(step["image"]).exists():
                ip=doc.add_paragraph();ip.alignment=1;ip.add_run().add_picture(step["image"],width=Cm(15.5));cp=doc.add_paragraph();cp.alignment=1;cp.add_run(step.get("caption",Path(step["image"]).name))
    out=Path(a.output);out.parent.mkdir(parents=True,exist_ok=True);doc.save(out);print(out)
if __name__=="__main__": main()
