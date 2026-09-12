#!/usr/bin/env python3
import argparse,csv,json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
def font(run,bold=False):
    run.font.name="微软雅黑";run._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑");run.font.size=Pt(10.5);run.bold=bold
def main():
    ap=argparse.ArgumentParser();ap.add_argument("project");ap.add_argument("output");a=ap.parse_args();root=Path(a.project);inv=json.loads((root/"台账/inventory.json").read_text(encoding="utf-8-sig"));
    with (root/"台账/functions.csv").open(encoding="utf-8-sig",newline="") as f: funcs=list(csv.DictReader(f))
    doc=Document();p=doc.add_paragraph();p.alignment=1;font(p.add_run(inv.get("project",root.name)+" 平台操作手册总览"),True)
    doc.add_heading("1 文档定位",1);doc.add_paragraph("本文件说明平台组成、角色、关键业务链路和独立手册索引；具体步骤以各子系统独立手册为准。")
    doc.add_heading("2 平台与子系统目录",1)
    for pdef in inv.get("platforms",[]):
        doc.add_heading(pdef.get("name","未命名平台"),2)
        for s in pdef.get("subsystems",[]): doc.add_paragraph(s.get("name","未命名子系统")+"（"+s.get("status","未验证")+"）",style="List Bullet")
    doc.add_heading("3 功能索引与关联",1);t=doc.add_table(rows=1,cols=7);t.style="Table Grid";heads=["编号","平台/子系统","功能","角色","上游","下游","章节"]
    for i,h in enumerate(heads):font(t.rows[0].cells[i].paragraphs[0].add_run(h),True)
    for x in funcs:
        vals=[x.get("function_id",""),x.get("platform","")+"/"+x.get("subsystem",""),x.get("menu","")+"/"+x.get("page",""),x.get("role",""),x.get("upstream",""),x.get("downstream",""),x.get("manual_chapter","")];cells=t.add_row().cells
        for i,v in enumerate(vals):font(cells[i].paragraphs[0].add_run(v))
    doc.add_heading("4 关键业务链路",1);doc.add_paragraph("依据台账/flows.csv 维护角色交接、状态变化和后台/小程序入口；不重复独立手册步骤。")
    out=Path(a.output);out.parent.mkdir(parents=True,exist_ok=True);doc.save(out);print(out)
if __name__=="__main__":main()
